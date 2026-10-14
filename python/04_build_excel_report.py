# STEP 4 - Build the Excel report
# Problem: Finance and managers don't open databases or code. They open Excel.
#          So the results have to end up in one clean workbook they can use right away.

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
# NOTE: pandas builds the tables. openpyxl = the tool that formats Excel files from Python
#       (colors, column widths, charts) - the stuff I'd normally do by hand.

flags = pd.read_csv("output/flagged_invoices.csv")
inv = pd.read_csv("output/invoices_clean.csv")
pricing = flags[flags["issue"] != "Likely duplicate"]
dups = flags[flags["issue"] == "Likely duplicate"]
# NOTE: Opened the flags from Step 3 and the clean invoices.
#       Split the flags into 2 groups: pricing problems (rate, quantity, fees) and duplicates.
#       I keep them apart because a duplicate is the WHOLE invoice at risk, not just the extra part.

summary = pd.DataFrame({
    "Metric": ["Invoices audited", "Invoices with a pricing issue", "Pricing overbilling ($)",
               "Likely duplicate invoices", "Duplicate dollars at risk ($)"],
    "Value": [len(inv), pricing["invoice_id"].nunique(), round(pricing["overbilled"].sum(), 2),
              len(dups), round(dups["overbilled"].sum(), 2)],
})
# NOTE: Built the Summary table - the 5 numbers a manager reads first.

by_vendor = (pricing.groupby("vendor_name")
             .agg(flagged_invoices=("invoice_id", "nunique"), overbilled=("overbilled", "sum"))
             .round(2).sort_values("overbilled", ascending=False).reset_index())
by_vendor["share_pct"] = (100 * by_vendor["overbilled"] / by_vendor["overbilled"].sum()).round(1)
# NOTE: groupby = a PivotTable. One row per vendor: how many invoices got flagged and how many dollars.
#       Sorted biggest first, then added each vendor's % of the total overbilling.

by_issue = (flags.groupby("issue")
            .agg(invoices=("invoice_id", "count"), dollars=("overbilled", "sum"))
            .round(2).reset_index())
# NOTE: Same idea, but one row per issue type - so I can see WHAT is causing the overbilling.

dispute = (flags.merge(inv[["invoice_id", "invoice_date", "amount_billed", "approval_status"]], on="invoice_id")
           .sort_values("overbilled", ascending=False))
# NOTE: The Dispute List - every flagged invoice with its date, amount and status, biggest first.
#       .merge = XLOOKUP: bring in the invoice details by matching invoice_id.
#       This is the sheet someone would actually work from to call vendors.

path = "output/Vendor_Billing_Audit.xlsx"
with pd.ExcelWriter(path) as xl:
    summary.to_excel(xl, sheet_name="Summary", index=False)
    by_vendor.to_excel(xl, sheet_name="By Vendor", index=False)
    by_issue.to_excel(xl, sheet_name="By Issue", index=False)
    dispute.to_excel(xl, sheet_name="Dispute List", index=False)
# NOTE: Wrote all 4 tables into ONE workbook - one sheet each.

wb = load_workbook(path)
for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for letter in ["A", "B", "C", "D", "E", "F", "G"]:
        ws.column_dimensions[letter].width = 24
# NOTE: Opened the workbook again with openpyxl and formatted every sheet:
#       header row (ws[1]) = bold white text on dark blue, and columns A-G set to width 24 so you can read them.

ws = wb["By Vendor"]
chart = BarChart()
chart.title = "Pricing overbilling by vendor ($)"
chart.type = "bar"
chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=ws.max_row), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
chart.height, chart.width = 9, 18
ws.add_chart(chart, "F2")
wb.save(path)
# NOTE: Added a real Excel bar chart on the By Vendor sheet - dollars (column C) by vendor name (column A).
#       Put it at cell F2 and saved the file.

print(summary.to_string(index=False))
print("\nSaved", path)
# NOTE: Printed the summary so I can see the headline numbers without opening Excel.
